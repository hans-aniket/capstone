from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_detailed_viva_report():
    doc = Document()
    
    # --- TITLE ---
    title = doc.add_heading('Comprehensive Project Report for Practical & Viva', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('Project: Comparative Sentiment Analysis on Product Reviews\n').alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # --- 1. INTRODUCTION ---
    doc.add_heading('1. Introduction & Objectives', level=1)
    doc.add_paragraph(
        "Sentiment analysis is a fundamental Natural Language Processing (NLP) task focused on determining the emotional tone behind a body of text. "
        "The objective of this capstone project is to systematically benchmark three distinct eras and tiers of Machine Learning approaches for binary sentiment classification:\n"
        "1. Classical Machine Learning (TF-IDF based)\n"
        "2. Deep Learning (Word Embedding based)\n"
        "3. Transformers (Attention based)\n\n"
        "This report details the methodology, system architecture, performance evaluation, and potential viva questions to defend the implementation."
    )

    # --- 2. DATASET ---
    doc.add_heading('2. Dataset & Preprocessing', level=1)
    doc.add_paragraph(
        "Dataset Used: The 'amazon_polarity' dataset sourced from HuggingFace. It consists of Amazon product reviews categorized into Positive (1) and Negative (0) sentiments.\n\n"
        "Data Subsampling: To optimize training times and manage computational resources without sacrificing statistical significance, "
        "a stratified subset was extracted using a fixed random seed (42).\n"
        "- Training Set: 20,000 samples\n"
        "- Testing Set: 5,000 samples\n\n"
        "Preprocessing Steps:\n"
        "- Classical ML: Lowercasing, punctuation removal, stop-word removal, and Vectorization using Term Frequency-Inverse Document Frequency (TF-IDF).\n"
        "- Deep Learning: Tokenization, sequence padding/truncation, and mapping to dense word embeddings (e.g., GloVe).\n"
        "- Transformers: Sub-word tokenization using the model-specific WordPiece tokenizer (e.g., DistilBERT tokenizer) generating input IDs and attention masks."
    )

    # --- 3. METHODOLOGY & MODELS ---
    doc.add_heading('3. Methodology & Model Architectures', level=1)
    
    doc.add_heading('3.1 Classical Machine Learning', level=2)
    doc.add_paragraph(
        "These models act as our robust baselines. They rely on sparse feature representations (TF-IDF).\n"
        "- Logistic Regression: A linear model estimating the probability of a binary response based on one or more predictor variables.\n"
        "- Support Vector Machine (Linear SVM): Finds the optimal hyperplane that maximizes the margin between positive and negative classes.\n"
        "- Multinomial Naive Bayes: A probabilistic classifier based on applying Bayes' theorem with strong independence assumptions between features."
    )

    doc.add_heading('3.2 Deep Learning', level=2)
    doc.add_paragraph(
        "These models capture sequential dependencies and local feature hierarchies using dense vectors.\n"
        "- TextCNN: Utilizes 1D Convolutional Neural Networks over word embeddings. Multiple filter sizes capture different n-gram features (e.g., bigrams, trigrams).\n"
        "- LSTM (Long Short-Term Memory): A specialized Recurrent Neural Network (RNN) designed to overcome the vanishing gradient problem, capturing long-term dependencies in the text.\n"
        "- CNN-LSTM Hybrid: Combines local feature extraction (CNN) with sequential pattern recognition (LSTM)."
    )

    doc.add_heading('3.3 Transformers', level=2)
    doc.add_paragraph(
        "Transformers rely on the self-attention mechanism to weigh the importance of all words in a sentence simultaneously.\n"
        "- DistilBERT: A distilled version of BERT (Bidirectional Encoder Representations from Transformers) that is smaller, faster, and lighter while retaining 97% of language understanding capabilities. Fine-tuned specifically on our dataset."
    )

    # --- 4. SYSTEM ARCHITECTURE (Practical Implementation) ---
    doc.add_heading('4. System Architecture & Deployment', level=1)
    doc.add_paragraph(
        "The project is structured as a full-stack web application designed for real-time inference.\n\n"
        "Backend (FastAPI):\n"
        "- Built in Python using the FastAPI framework for high performance.\n"
        "- Model weights are pre-loaded into memory during the application lifecycle (lifespan events) to prevent blocking on inference requests.\n"
        "- Handles Cross-Origin Resource Sharing (CORS) and serves static files.\n\n"
        "Frontend (React & Vite):\n"
        "- A modern, dynamic user interface built with React.\n"
        "- Communicates with the FastAPI backend via REST endpoints to display comparative metrics and allow users to test custom text inputs.\n\n"
        "Deployment (Render):\n"
        "- Infrastructure as Code (IaC) is managed via 'render.yaml'.\n"
        "- Combines both frontend and backend into a single Web Service for cost-efficiency on Render's cloud platform."
    )

    # --- 5. EVALUATION METRICS & RESULTS ---
    doc.add_heading('5. Evaluation Metrics & Experimental Results', level=1)
    doc.add_paragraph(
        "The models were evaluated primarily on Classification Accuracy, ROC-AUC (Area Under the Receiver Operating Characteristic Curve), and Inference Latency (milliseconds per sample)."
    )

    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Model'
    hdr_cells[1].text = 'Accuracy'
    hdr_cells[2].text = 'ROC-AUC'
    hdr_cells[3].text = 'Inference Latency (ms)'
    
    results = [
        ("SVM (Linear)", "90.16%", "0.966", "0.003 ms"),
        ("Logistic Regression", "89.02%", "0.958", "0.001 ms"),
        ("Naive Bayes", "88.06%", "0.950", "0.003 ms"),
        ("TextCNN", "85.64%", "0.933", "1.467 ms"),
        ("CNN-LSTM", "85.12%", "0.927", "1.372 ms"),
        ("LSTM", "76.60%", "0.861", "10.952 ms"),
        ("DistilBERT*", "50.00%", "0.482", "94.820 ms")
    ]
    
    for model, acc, auc, latency in results:
        row_cells = table.add_row().cells
        row_cells[0].text = model
        row_cells[1].text = acc
        row_cells[2].text = auc
        row_cells[3].text = latency
        
    doc.add_paragraph(
        "\n*Note: DistilBERT was evaluated on a dramatically reduced subset (n=40) during initial testing and exhibited baseline/random performance, indicating an under-trained state compared to the robustly trained classical and DL models."
    )

    doc.add_heading('5.1 Conclusion', level=2)
    doc.add_paragraph(
        "Classical Models Excel: Traditional models leveraging TF-IDF proved highly effective for sentiment polarity on Amazon reviews, achieving the highest accuracy (Linear SVM at 90.16%) with near-zero latency.\n"
        "Deep Learning Overhead: TextCNN and LSTM architectures incurred significantly higher computational costs (both training and inference) without surpassing the baseline TF-IDF approaches, suggesting that simple lexical features are often sufficient for basic polarity classification.\n"
        "Practical Trade-offs: For real-time production systems, Logistic Regression or Linear SVM are the strongest candidates due to their efficiency."
    )

    # --- 6. VIVA VOCE Q&A ---
    doc.add_heading('6. Potential Viva Voce Questions & Answers', level=1)

    qa_pairs = [
        ("Q1: Why did Classical ML (SVM/Logistic Regression) outperform Deep Learning models here?",
         "A1: Sentiment analysis on product reviews often relies heavily on specific strong keywords (e.g., 'excellent', 'terrible'). TF-IDF explicitly highlights these important words. Deep Learning models require massive amounts of data to learn semantic relationships from scratch; our 20k sample size might not have been large enough for LSTMs to fully realize their potential advantage over highly optimized classical baselines."),
        
        ("Q2: What is TF-IDF and how does it work?",
         "A2: Term Frequency-Inverse Document Frequency evaluates how relevant a word is to a document in a collection. 'Term Frequency' counts word occurrences in a document, while 'Inverse Document Frequency' penalizes words that appear frequently across all documents (like 'the', 'is'), giving higher weight to unique, meaningful words."),
        
        ("Q3: Explain the difference between CNNs and LSTMs for text.",
         "A3: TextCNNs use convolutional filters to identify local patterns or n-grams (e.g., short phrases) regardless of their position in the text. LSTMs process text sequentially, maintaining a 'hidden state' memory that allows them to understand long-term dependencies and context across a longer sentence."),
        
        ("Q4: Why was FastAPI chosen over Flask or Django for the backend?",
         "A4: FastAPI is extremely fast due to its asynchronous capabilities (ASGI) and Pydantic-based data validation. It also natively supports lifespan events, which is crucial for loading heavy machine learning models into memory just once upon server startup, rather than blocking individual requests."),
        
        ("Q5: What challenges did you face with the Transformer model (DistilBERT)?",
         "A5: The primary challenge is computational overhead. Transformers require significant GPU memory for fine-tuning. As seen in the results, DistilBERT requires careful hyperparameter tuning and sufficient data epochs; a lack of these (resulting in 50% accuracy on a tiny test set) shows that Transformers are not 'plug-and-play' and require heavy computational investment compared to classical models."),
         
        ("Q6: How did you deploy the application?",
         "A6: The application was containerized/deployed on Render as a single Web Service. The React frontend is built as static assets (`npm run build`) and the FastAPI backend serves these static files via its root endpoint, while handling API inference requests via specific `/api` routes. Dependencies were managed via `requirements.txt`.")
    ]

    for q, a in qa_pairs:
        p_q = doc.add_paragraph()
        p_q.add_run(q).bold = True
        p_a = doc.add_paragraph(a)
        p_a.style = 'Body Text'

    doc.save('Capstone_Viva_Practical_Report.docx')

if __name__ == "__main__":
    create_detailed_viva_report()
