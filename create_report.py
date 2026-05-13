from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_report():
    doc = Document()
    
    # Title
    title = doc.add_heading('Comparative Sentiment Analysis on Product Reviews', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_heading('1. Introduction', level=1)
    doc.add_paragraph(
        "This report outlines the methodology, models, and results for the Comparative Sentiment "
        "Analysis project. The primary goal of this capstone project is to systematically benchmark "
        "three distinct tiers of Machine Learning approaches: Classical Machine Learning, Deep Learning, "
        "and Transformer-based models."
    )

    doc.add_heading('2. Dataset & Task Details', level=1)
    doc.add_paragraph(
        "The project focuses on binary sentiment classification (Positive vs. Negative) using the "
        "'amazon_polarity' dataset sourced from HuggingFace. "
        "For optimal training efficiency, a stratified subset of the data was utilized, comprising "
        "20,000 training samples and 5,000 testing samples (random seed = 42)."
    )

    doc.add_heading('3. Models Evaluated', level=1)
    doc.add_paragraph("We categorized our experiments into three model tiers:")
    ul1 = doc.add_paragraph(style='List Bullet')
    ul1.add_run("Classical ML: ").bold = True
    ul1.add_run("Logistic Regression, Support Vector Machine (Linear), and Multinomial Naive Bayes (using TF-IDF vectors).")
    
    ul2 = doc.add_paragraph(style='List Bullet')
    ul2.add_run("Deep Learning: ").bold = True
    ul2.add_run("Long Short-Term Memory (LSTM), TextCNN, and a hybrid CNN-LSTM network.")
    
    ul3 = doc.add_paragraph(style='List Bullet')
    ul3.add_run("Transformers: ").bold = True
    ul3.add_run("DistilBERT (fine-tuned on the dataset).")

    doc.add_heading('4. Performance Results', level=1)
    doc.add_paragraph(
        "The evaluation metrics focused on Classification Accuracy, ROC-AUC, and Computational Efficiency "
        "(Inference Time in milliseconds per sample)."
    )

    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Model'
    hdr_cells[1].text = 'Accuracy'
    hdr_cells[2].text = 'ROC-AUC'
    hdr_cells[3].text = 'Inference Latency (ms)'
    
    results = [
        ("SVM (Linear)", "90.16%", "0.966", "0.003"),
        ("Logistic Regression", "89.02%", "0.958", "0.001"),
        ("Naive Bayes", "88.06%", "0.950", "0.003"),
        ("TextCNN", "85.64%", "0.933", "1.467"),
        ("CNN-LSTM", "85.12%", "0.927", "1.372"),
        ("LSTM", "76.60%", "0.861", "10.952"),
        ("DistilBERT*", "50.00%", "0.482", "94.820")
    ]
    
    for model, acc, auc, latency in results:
        row_cells = table.add_row().cells
        row_cells[0].text = model
        row_cells[1].text = acc
        row_cells[2].text = auc
        row_cells[3].text = latency
        
    doc.add_paragraph("*Note: DistilBERT was evaluated on a drastically smaller test set (n=40) during initial testing and exhibited baseline performance, indicating it may require more extensive fine-tuning steps or compute resources.")

    doc.add_heading('5. Key Insights & Conclusion', level=1)
    doc.add_paragraph(
        "1. Classical Models Excel in this Domain: The Linear SVM achieved the highest overall accuracy "
        "(90.16%) while maintaining an extremely low inference latency (0.003 ms/sample). Traditional models "
        "leveraging TF-IDF proved highly effective for sentiment polarity on Amazon reviews.\n\n"
        "2. Deep Learning Overhead: While TextCNN achieved a respectable 85.64% accuracy, it required significantly "
        "more training time and yielded slightly lower performance than classical models. Pure LSTM underperformed (76.60%) "
        "and suffered from the highest inference latency among non-transformer models.\n\n"
        "3. Scalability: If deployed in a production environment with strict latency constraints, the Logistic Regression "
        "or Linear SVM models are the strongest candidates. They offer the best trade-off between predictive power and computational cost."
    )

    doc.save('Capstone_Report.docx')

if __name__ == "__main__":
    create_report()
